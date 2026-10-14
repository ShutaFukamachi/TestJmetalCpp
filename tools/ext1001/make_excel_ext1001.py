#!/usr/bin/env python3
"""
make_excel_ext1001.py -- 2026-10-01 ゼミ用の数値比較ブック（拡張コスト表）

入力: cmake-build-release/analysis/ext1001/ の集計 CSV / JSON（run_all.py が作る）
出力: cmake-build-release/analysis/ext1001/EXT1001_比較データ.xlsx
"""
import csv, json, os, statistics as S
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
BD = os.path.normpath(os.path.join(HERE, "..", "..", "cmake-build-release"))
A = os.path.join(BD, "analysis", "ext1001")
M3 = ["SchedObj", "MaxShift", "MaxShiftDur"]
CONDS = ["RR000_RV0", "RR000_RV1", "RR025_RV0", "RR025_RV1", "RR050_RV0", "RR050_RV1"]
HEAD = PatternFill("solid", fgColor="DDE7F5")


def rd(name):
    p = os.path.join(A, name)
    return list(csv.DictReader(open(p, encoding="utf-8"))) if os.path.exists(p) else []


def J(name):
    p = os.path.join(A, name)
    return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else {}


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def sheet(wb, title, header, rows, note=None, widths=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        for line in note.split("\n"):
            ws.cell(row=r0, column=1, value=line).font = Font(italic=True, color="555555")
            r0 += 1
        r0 += 1
    for j, h in enumerate(header, 1):
        c = ws.cell(row=r0, column=j, value=h); c.font = Font(bold=True); c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, float):
                c.number_format = "0.0000" if abs(v) < 10 else "#,##0.0"
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    for j in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, 14)
    return ws


def main():
    s = J("summary.json"); rho = J("rho_summary.json"); art = J("artifact_summary.json")
    rows = rd("rows.csv"); feat = {r["instance"]: r for r in rd("features.csv")}
    wb = Workbook(); wb.remove(wb.active)

    # ---- 概要 ----
    hv = s.get("hv_vs_mip_same_subset", {}); t30 = s.get("three_way", {}).get("j30", {}); t60 = s.get("three_way", {}).get("j60", {})
    rp = s.get("repro", {}); e1 = s.get("sweep_E1", {}); e2 = s.get("sweep_E2", {})
    ov = []
    ov.append(["指摘1 コスト表拡張", "旧MIP前線のうちコスト表範囲外を使う点の割合", art.get("frac_uses_clamp")])
    ov.append(["指摘1 コスト表拡張", "その点を正しいコストで再計算した増加率（平均）", art.get("clamp_mean_rel_change")])
    ov.append(["指摘1 コスト表拡張", "同（最大）", art.get("clamp_max_rel_change")])
    ov.append(["指摘1 コスト表拡張", "拡張MIP前線に負ける旧点 / 範囲外の旧点", f"{art.get('clamp_dominated')}/{art.get('n_clamp')}"])
    for m in ["MaxShiftDur", "MaxShift", "SchedObj"]:
        e, o = hv.get(f"ext_{m}", {}), hv.get(f"old_{m}", {})
        ov.append(["指摘2 MIP比HV（拡張）", m, e.get("hv_ratio")])
        ov.append(["指摘2 MIP比HV（旧・同じ条件）", m, o.get("hv_ratio")])
        ov.append(["指摘2 MIP前線に到達した割合", m, e.get("reach")])
    ov.append(["指摘2 3手法比較 j30", "条件数 / インスタンス数", f"{t30.get('n')} / {t30.get('n_inst')}"])
    ov.append(["指摘2 3手法比較 j30", "MaxShiftDur > SchedObj の条件数", t30.get("msd_beats_so")])
    ov.append(["指摘2 3手法比較 j30", "SchedObj > MaxShiftDur の条件数", t30.get("so_beats_msd")])
    ov.append(["指摘4 3手法比較 j60", "条件数 / インスタンス数", f"{t60.get('n')} / {t60.get('n_inst')}"])
    ov.append(["指摘4 3手法比較 j60", "SchedObj > MaxShiftDur の条件数", t60.get("so_beats_msd")])
    ov.append(["指摘3 max_shift感度", "1活動の窓幅を動かしたときのスケジュール変化回数（中央値）", e2.get("median_changes")])
    ov.append(["指摘3 max_shift感度", "MIPと同じスケジュールになる窓幅の範囲（中央値）", e2.get("median_frac_same")])
    ov.append(["指摘3 max_shift感度", "広げるとコストが悪化する変化を含む活動の割合", e2.get("frac_jobs_cost_up")])
    ov.append(["指摘3 max_shift感度", "全活動一緒: コスト最小の窓幅が途中にある掃引の割合", e1.get("frac_costmin_interior")])
    ov.append(["指摘5 再現性", "MIP解まるごと再現 MaxShiftDur", rp.get("sched_msd")])
    ov.append(["指摘5 再現性", "MIP解まるごと再現 MaxShift", rp.get("sched_ms")])
    ov.append(["指摘5 再現性", "MIP解まるごと再現 SchedObj（確実）", rp.get("sched_so_det")])
    ov.append(["指摘5 再現性", "活動単位 MaxShiftDur / SchedObj（確実）", f"{rp.get('job_msd',0):.4f} / {rp.get('job_so_det',0):.4f}"])
    for k, lab in [("noise_rp1_minus_base", "k=1同士の差（ノイズ）"), ("rp3_minus_k1avg", "k=3 − k=1"),
                   ("rp5_minus_k1avg", "k=5 − k=1"), ("rp1q_minus_k1avg", "評価1/4 k=1 − 評価フル k=1"),
                   ("rp3q_minus_rp1q", "評価1/4: k=3 − k=1")]:
        v = rho.get(k) or {}
        ov.append(["指摘6 ρべき乗", lab + "（共通HVの差の平均）", v.get("mean")])
    sheet(wb, "概要", ["指摘", "指標", "値"], ov,
          note="2026-10-01 ゼミ用。拡張コスト表（costs_ext）での再計算結果の要約。\n"
               "HV比は公式定義（MIP∪手法で正規化、参照点(1.1,1.1)）。割合は 0〜1 表記。",
          widths={1: 26, 2: 56, 3: 16})

    # ---- 条件別 MIP比 ----
    by = defaultdict(dict)
    for r in rows:
        if r["size"] == "j30" and r["hv_ratio"] != "":
            by[(r["instance"], r["cond"])][(r["root"], r["method"])] = r
    out = []
    for (inst, cond), d in sorted(by.items()):
        if not any(k[0] == "ext" for k in d):
            continue
        row = [inst, cond]
        for m in M3:
            e = d.get(("ext", m), {}); o = d.get(("old", m), {})
            row += [num(e.get("hv_ratio")), num(o.get("hv_ratio")), num(e.get("cost_gap")), num(e.get("ms_gap")), num(e.get("reach"))]
        out.append(row)
    hdr = ["instance", "cond"]
    for m in M3:
        hdr += [f"{m} HV比(拡張)", f"{m} HV比(旧)", f"{m} コストgap", f"{m} 工期gap", f"{m} 到達率"]
    sheet(wb, "MIP比_条件別", hdr, out, note="拡張MIPとNSGAの両方がそろった条件のみ。gap は割合（0.01=1%）。到達率=MIP前線の点のうちNSGAが弱支配した割合。")

    # ---- 3手法比較 ----
    tw = defaultdict(dict)
    for r in rows:
        if r["root"] == "ext":
            tw[(r["size"], r["instance"], r["cond"])][r["method"]] = num(r["hv_common"])
    out = []
    for (size, inst, cond), d in sorted(tw.items()):
        if len(d) < 3:
            continue
        best = max(d, key=d.get)
        f = feat.get(inst, {})
        out.append([size, inst, cond, d["SchedObj"], d["MaxShift"], d["MaxShiftDur"], d["MaxShiftDur"] - d["SchedObj"],
                    best, num(f.get("NC")), num(f.get("RF")), num(f.get("RS"))])
    sheet(wb, "3手法比較", ["size", "instance", "cond", "SchedObj 共通HV", "MaxShift 共通HV", "MaxShiftDur 共通HV",
                         "差 MSD−SO", "最良", "NC", "RF", "RS"], out,
          note="共通HV: その条件の全手法（MIPがあれば含む）の点をまとめて正規化したHV。手法間で直接比べられる。")

    # ---- 特徴量別 ----
    out = []
    for key, d in s.get("by_feature_j30", {}).items():
        for val, v in d.items():
            out.append([key, val, v.get("n"), v.get("mean_diff"), v.get("so_wins")])
    for c, v in t30.get("by_cond", {}).items():
        out.append(["条件", c, v.get("n"), v.get("mean_diff"), v.get("so_wins")])
    sheet(wb, "特徴量別", ["特徴", "値", "条件数", "差 MSD−SO（平均）", "SchedObjが勝った条件数"], out,
          note="j30。NC=ネットワークの複雑さ、RF=資源の使われ方、RS=資源の余裕（PSPLIBの設計値）。")

    # ---- 再現性 ----
    sch = rd("repro_sched.csv")
    agg = defaultdict(list)
    for r in sch:
        agg[(r["instance"], r["cond"])].append(r)
    out = []
    for (inst, cond), L in sorted(agg.items()):
        lp = sorted(float(x["so_log10p"]) for x in L)
        out.append([inst, cond, len(L), sum(int(x["msd_all"]) for x in L) / len(L), sum(int(x["ms_all"]) for x in L) / len(L),
                    sum(int(x["so_det_all"]) for x in L) / len(L), lp[len(lp) // 2]])
    sheet(wb, "再現性", ["instance", "cond", "MIP解の数", "MaxShiftDur 全体再現率", "MaxShift 全体再現率",
                       "SchedObj 確実な全体再現率", "SchedObj 運で全体再現する確率 log10（中央値）"], out,
          note="拡張MIP解を正解とし、各エンコーディングのデコーダ規則で全活動を同じ時刻に置けるかを厳密判定。")

    # ---- max_shift 感度 ----
    e1r = rd("sweep_E1_MIP.csv")
    out = [[r["instance"], r["cond"], int(r["front_pos"]), int(r["Wmax"]), int(r["n_distinct"]), int(r["n_change"]),
            int(r["max_plateau"]) if r["max_plateau"] else None, int(r["cost_up_changes"]), int(r["both_worse_changes"]),
            int(r["w_at_cost_min"]), num(r["cost_w0"]), num(r["cost_min"]), num(r["cost_wmax"]),
            int(float(r["ms_w0"])), int(float(r["ms_at_cost_min"])), int(r["n_nondominated_w"])] for r in e1r]
    sheet(wb, "maxshift感度_全活動", ["instance", "cond", "前線位置", "窓幅上限", "異なるスケジュール数", "変化回数", "最長の平ら区間",
                                  "広げてコスト悪化した回数", "工期もコストも悪化した回数", "コスト最小の窓幅",
                                  "コスト(w=0)", "コスト最小", "コスト(w=上限)", "工期(w=0)", "工期(コスト最小時)", "非劣な窓幅の数"], out,
          note="活動リスト固定・全活動に同じ窓幅 w を 0〜上限で1刻み。基準はMIP前線の工期最小端/中間/コスト最小端。")
    e2r = rd("sweep_E2_MIP.csv")
    g = defaultdict(list)
    for r in e2r:
        g[(r["instance"], r["cond"])].append(r)
    out = [[i, c, len(L), S.median(int(x["n_change"]) for x in L), S.median(float(x["frac_same_as_base"]) for x in L),
            sum(1 for x in L if int(x["cost_up"]) > 0) / len(L), sum(1 for x in L if int(x["n_change"]) == 0) / len(L)]
           for (i, c), L in sorted(g.items())]
    sheet(wb, "maxshift感度_1活動", ["instance", "cond", "掃引数", "変化回数（中央値）", "MIPと同じになる窓幅の範囲（中央値）",
                                  "コスト悪化の変化を含む割合", "一度も変わらない割合"], out,
          note="MIP解を再現する窓幅から、1活動だけ窓幅を 0〜上限で1刻みに動かす（遺伝子1つを変える状況）。")

    # ---- ρ 実験 ----
    out = []
    for k, v in rho.items():
        if isinstance(v, dict) and "mean" in v:
            out.append([k, v.get("n"), v.get("mean"), v.get("sd"), v.get("se"), v.get("pos"), v.get("neg")])
    for c, v in rho.get("by_cond", {}).items():
        out.append([f"条件別 {c}", v.get("n"), v.get("rp3"), None, None, "k=3", None])
        out.append([f"条件別 {c}", v.get("n"), v.get("rp5"), None, None, "k=5", None])
    sheet(wb, "rho実験", ["比較", "条件数", "平均差", "標準偏差", "標準誤差", "改善", "悪化"], out,
          note="窓幅 win=round(ρ^k×上限)。共通HVの差。k1avg=公式設定の実行とk=1再実行の平均。q=評価回数1/4。")

    # ---- コスト表拡張の影響 ----
    ar = rd("artifact_rows.csv")
    g = defaultdict(list)
    for r in ar:
        g[(r["instance"], r["cond"])].append(r)
    out = []
    for (i, c), L in sorted(g.items()):
        cl = [x for x in L if x["uses_clamp"] == "1"]
        out.append([i, c, len(L), len(cl), (sum(float(x["rel_change"]) for x in cl) / len(cl)) if cl else None,
                    max((float(x["rel_change"]) for x in cl), default=None), sum(int(x["dominated_by_ext"]) for x in cl)])
    sheet(wb, "コスト表拡張の影響", ["instance", "cond", "旧MIP前線の点数", "範囲外を使う点", "再計算の増加率（平均）",
                                "同（最大）", "拡張MIP前線に負ける点"], out,
          note="旧コスト表（t>=T0 を最終値に固定）の MIP 前線の各点を、拡張コスト表で計算し直した結果。")

    # ---- P1/P2/P3 ----
    E = os.path.join(BD, "results", "EXT_COST")
    srcs = [("P1", "RR000_RV0", os.path.join(E, "FUN_MIP", "j309_1", "GAP_MIP_j309_1_RR000_RV0.csv")),
            ("P3", "RR000_RV0", os.path.join(E, "FUN_MIP_P3SLOT", "j309_1", "GAP_MIP_P3SLOT_j309_1_RR000_RV0.csv")),
            ("P1", "RR000_RV1", os.path.join(E, "FUN_MIP", "j309_1", "GAP_MIP_j309_1_RR000_RV1.csv")),
            ("P2", "RR000_RV1", os.path.join(E, "FUN_MIP_P2", "j309_1", "GAP_MIP_P2_j309_1_RR000_RV1.csv")),
            ("P3", "RR000_RV1", os.path.join(E, "FUN_MIP_P3SLOT", "j309_1", "GAP_MIP_P3SLOT_j309_1_RR000_RV1.csv"))]
    out = []
    for mode, cond, p in srcs:
        for r in (list(csv.DictReader(open(p))) if os.path.exists(p) else []):
            out.append(["j309_1", cond, mode, num(r["makespan"]), num(r["cost"]), num(r["gap"]), r["optimal"], num(r["solve_seconds"])])
    sheet(wb, "P1_P2_P3", ["instance", "cond", "モード", "メイクスパン", "コスト", "gap", "最適性証明", "求解秒"], out,
          note="P1=分割なし、P2=休暇時のみ中断、P3=任意中断（新定式化 P3Slot）。すべて拡張コスト表。optimal=0 は時間切れ（120秒）。")

    path = os.path.join(A, "EXT1001_比較データ.xlsx")
    wb.save(path)
    print("[excel]", path)


if __name__ == "__main__":
    main()
